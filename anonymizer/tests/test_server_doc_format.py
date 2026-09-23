"""Старый .doc возвращается документом Word, а не .txt.

Записать обезличенную копию обратно в бинарный Word 97-2003 нечем, но и
отдавать её текстом нельзя: .txt рвёт документ окончательно, и восстанавливать
по маппингу становится нечего (см. _WORD_EXTENSIONS в server.py). Поэтому .doc
поднимается до .docx — через LibreOffice, когда она есть на хосте (разметка
сохраняется), и простым документом из текста, когда её нет.

Пайплайн здесь минимальный (все стадии детекции выключены, как в
test_server_auth.py): проверяется маршрутизация форматов, а не детекторы.
LibreOffice в тестах не запускается — её наличие/отсутствие подменяется на
``documents.doc_to_docx_bytes``.
"""

from __future__ import annotations

import base64
import io
import json
import sys
import tempfile
from contextlib import contextmanager
from pathlib import Path

import docx
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from anonymizer import documents, server, usage_log  # noqa: E402

_TEXT = "Первый абзац документа.\nВторой абзац документа."
# Сигнатура OLE2 — содержимое неважно: чтение .doc в тестах подменяется.
_FAKE_DOC = b"\xd0\xcf\x11\xe0 fake ole"


@contextmanager
def _minimal_pipeline():
    orig = (
        server._DETECTORS,
        server._DEFAULTS,
        server._REVIEW_CFG,
        server._NER_BACKEND,
        server._NEEDS_MODEL_LOCK,
    )
    server._DETECTORS = {}
    server._DEFAULTS = {name: False for name in server._STAGE_NAMES}
    server._REVIEW_CFG = None
    server._NER_BACKEND = "none"
    server._NEEDS_MODEL_LOCK = False
    try:
        yield
    finally:
        (
            server._DETECTORS,
            server._DEFAULTS,
            server._REVIEW_CFG,
            server._NER_BACKEND,
            server._NEEDS_MODEL_LOCK,
        ) = orig


@contextmanager
def _temp_usage_log():
    orig = usage_log.LOG_PATH
    with tempfile.TemporaryDirectory() as tmp:
        usage_log.LOG_PATH = Path(tmp) / "usage.jsonl"
        try:
            yield
        finally:
            usage_log.LOG_PATH = orig


def _docx_bytes(text: str) -> bytes:
    document = docx.Document()
    for line in text.split("\n"):
        document.add_paragraph(line)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def _docx_text(data: bytes) -> str:
    return "\n".join(p.text for p in docx.Document(io.BytesIO(data)).paragraphs)


@pytest.fixture
def no_libreoffice(monkeypatch):
    """Хост без LibreOffice: конвертер молча недоступен."""
    monkeypatch.setattr(documents, "doc_to_docx_bytes", lambda data, suffix=".doc": None)
    # Текст из .doc в этом случае достаёт antiword/catdoc/olefile — их в
    # тестовом окружении нет, поэтому подменяем сам извлекатель.
    monkeypatch.setattr(documents, "_read_doc_bytes", lambda data: _TEXT)


@pytest.fixture
def with_libreoffice(monkeypatch):
    """Хост с LibreOffice: .doc конвертируется в настоящий .docx."""
    monkeypatch.setattr(
        documents, "doc_to_docx_bytes", lambda data, suffix=".doc": _docx_bytes(_TEXT)
    )


# --- Мини-каркас HTTP-обработчика без сокета (как в test_server_auth.py) ---

class _Headers:
    def __init__(self, items: dict | None = None) -> None:
        self._items = {k.lower(): v for k, v in (items or {}).items()}

    def get(self, key: str, default=None):
        return self._items.get(key.lower(), default)


def _post_json(path: str, body_obj) -> tuple[int, dict]:
    body = json.dumps(body_obj).encode("utf-8")
    h = server.Handler.__new__(server.Handler)
    h.command = "POST"
    h.path = path
    h.request_version = "HTTP/1.1"
    h.protocol_version = "HTTP/1.1"
    h.requestline = f"POST {path} HTTP/1.1"
    h.client_address = ("127.0.0.1", 55555)
    h.headers = _Headers({"Content-Length": str(len(body))})
    h.rfile = io.BytesIO(body)
    h.wfile = io.BytesIO()
    h.close_connection = True
    h.do_POST()
    raw = h.wfile.getvalue()
    head, _, resp = raw.partition(b"\r\n\r\n")
    return int(head.split(b"\r\n", 1)[0].split()[1]), (json.loads(resp) if resp else {})


@contextmanager
def _anonymous_access():
    orig_keys, orig_allow = server._API_KEYS, server._ALLOW_ANONYMOUS
    server._API_KEYS, server._ALLOW_ANONYMOUS = {}, True
    try:
        yield
    finally:
        server._API_KEYS, server._ALLOW_ANONYMOUS = orig_keys, orig_allow


def _anonymize(filename: str, raw: bytes) -> dict:
    with _minimal_pipeline(), _temp_usage_log():
        return server._run_anonymize_file(
            {"filename": filename, "file_base64": base64.b64encode(raw).decode("ascii")}
        )


def test_doc_without_libreoffice_still_comes_back_as_word(no_libreoffice):
    res = _anonymize("договор.doc", _FAKE_DOC)
    assert res["document_name"] == "договор.anon.docx"
    assert res["document_mime"] == server._DOCX_MIME
    assert res["is_docx"] is True
    # Разметки в нём нет, и клиенту об этом сказано честно.
    assert res["document_source"] == "text"
    # Файл действительно открывается как Word-документ, а не подписан .docx.
    assert _docx_text(base64.b64decode(res["document_base64"])) == _TEXT


def test_doc_with_libreoffice_goes_through_the_structure_preserving_path(with_libreoffice):
    res = _anonymize("договор.doc", _FAKE_DOC)
    assert res["document_name"] == "договор.anon.docx"
    assert res["is_docx"] is True
    assert res["document_source"] == "converted"
    assert _docx_text(base64.b64decode(res["document_base64"])) == _TEXT
    # Текст берётся из конвертированного документа — иначе цель переписывания
    # и источник текста разъезжаются (см. комментарий в _run_anonymize_file).
    assert res["anonymized_text"] == _TEXT


def test_docx_path_is_unchanged(with_libreoffice):
    res = _anonymize("договор.docx", _docx_bytes(_TEXT))
    assert res["document_name"] == "договор.anon.docx"
    assert res["is_docx"] is True
    assert res["document_source"] == "original"
    assert _docx_text(base64.b64decode(res["document_base64"])) == _TEXT


def test_other_formats_still_come_back_as_text():
    res = _anonymize("заметка.txt", _TEXT.encode("utf-8"))
    assert res["document_name"] == "заметка.anon.txt"
    assert res["is_docx"] is False
    assert res["document_mime"] == "text/plain"


def test_restoring_a_doc_gives_back_a_word_document(monkeypatch):
    """Обратная сторона той же задачи: восстановление .doc тоже не должно
    вырождаться в .txt — иначе восстановленный документ снова приходится
    собирать вручную."""
    masked = "Первый абзац [PERSON_1].\nВторой абзац документа."
    monkeypatch.setattr(
        documents, "doc_to_docx_bytes", lambda data, suffix=".doc": _docx_bytes(masked)
    )
    with _anonymous_access(), _temp_usage_log():
        code, res = _post_json(
            "/deanonymize-file",
            {
                "filename": "договор.anon.doc",
                "file_base64": base64.b64encode(_FAKE_DOC).decode("ascii"),
                "mapping": {"[PERSON_1]": "Иванова"},
            },
        )
    assert code == 200, res
    assert res["document_name"] == "договор.restored.docx"
    assert res["document_mime"] == server._DOCX_MIME
    assert res["is_docx"] is True
    restored = _docx_text(base64.b64decode(res["document_base64"]))
    assert "Иванова" in restored and "[PERSON_1]" not in restored


def test_restoring_a_doc_without_libreoffice_is_still_a_word_document(no_libreoffice):
    with _anonymous_access(), _temp_usage_log():
        code, res = _post_json(
            "/deanonymize-file",
            {
                "filename": "договор.anon.doc",
                "file_base64": base64.b64encode(_FAKE_DOC).decode("ascii"),
                "mapping": {"[PERSON_1]": "Иванова"},
            },
        )
    assert code == 200, res
    assert res["document_name"] == "договор.restored.docx"
    assert res["is_docx"] is True
    # Разметки взять неоткуда, но это документ Word с восстановленным текстом.
    assert _docx_text(base64.b64decode(res["document_base64"])) == _TEXT


def test_unreadable_conversion_falls_back_instead_of_failing(monkeypatch):
    """LibreOffice может отдать файл, который python-docx не читает. Это не
    повод уронить запрос: откатываемся на прежний путь."""
    monkeypatch.setattr(
        documents, "doc_to_docx_bytes", lambda data, suffix=".doc": b"not a docx at all"
    )
    monkeypatch.setattr(documents, "_read_doc_bytes", lambda data: _TEXT)
    res = _anonymize("договор.doc", _FAKE_DOC)
    assert res["document_name"] == "договор.anon.docx"
    assert res["document_source"] == "text"
    assert _docx_text(base64.b64decode(res["document_base64"])) == _TEXT
