"""PDF возвращается PDF-ом: pdf → .odg (LibreOffice Draw) → правка → pdf.

Скан (PDF без текстового слоя) отклоняется, а при отсутствии LibreOffice
.doc/.xls/.rtf/.pdf всё равно отдаются, но с предупреждением о потере разметки.

Тесты с реальной LibreOffice пропускаются, если её на хосте нет. Остальные
LibreOffice не запускают.
"""

from __future__ import annotations

import base64
import io
import struct
import sys
import tempfile
import zipfile
import zlib
from contextlib import contextmanager
from pathlib import Path

import docx
import pytest
from docx.shared import Inches

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from anonymizer import documents, server, usage_log  # noqa: E402
from anonymizer.detectors import DEFAULT_DETECTORS  # noqa: E402

_NEEDS_SOFFICE = pytest.mark.skipif(
    documents._find_soffice() is None, reason="LibreOffice не установлена"
)
_EMAIL = "ivan.petrov@example.ru"
_INN = "380101234567"


@contextmanager
def _regex_pipeline():
    orig = (
        server._DETECTORS,
        server._DEFAULTS,
        server._REVIEW_CFG,
        server._NER_BACKEND,
        server._NEEDS_MODEL_LOCK,
    )
    server._DETECTORS = {"regex": list(DEFAULT_DETECTORS)}
    server._DEFAULTS = {name: False for name in server._STAGE_NAMES}
    server._DEFAULTS["regex"] = True
    server._REVIEW_CFG = None
    server._NER_BACKEND = "none"
    server._NEEDS_MODEL_LOCK = False
    try:
        yield
    finally:
        (
            server._DETECTORS,
            server._DEFAULTS,
            server._REVIEW_CFG,
            server._NER_BACKEND,
            server._NEEDS_MODEL_LOCK,
        ) = orig


@contextmanager
def _temp_usage_log():
    orig = usage_log.LOG_PATH
    with tempfile.TemporaryDirectory() as tmp:
        usage_log.LOG_PATH = Path(tmp) / "usage.jsonl"
        try:
            yield
        finally:
            usage_log.LOG_PATH = orig


def _anonymize(filename: str, raw: bytes) -> dict:
    with _regex_pipeline(), _temp_usage_log():
        return server._run_anonymize_file(
            {"filename": filename, "file_base64": base64.b64encode(raw).decode("ascii")}
        )


def _png(size: int = 40) -> bytes:
    def chunk(tag: bytes, body: bytes) -> bytes:
        return (
            struct.pack(">I", len(body))
            + tag
            + body
            + struct.pack(">I", zlib.crc32(tag + body) & 0xFFFFFFFF)
        )

    raw = b"".join(b"\x00" + b"\xff\x00\x00" * size for _ in range(size))
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", size, size, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )


def _docx_with_image() -> bytes:
    document = docx.Document()
    document.add_paragraph(f"Заказчик Иванов Иван Иванович, ИНН {_INN}.")
    document.add_picture(io.BytesIO(_png()), width=Inches(1))
    document.add_paragraph(f"Контакт {_EMAIL}")
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def _docx_only_image() -> bytes:
    document = docx.Document()
    document.add_picture(io.BytesIO(_png()), width=Inches(4))
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def _to_pdf(docx_bytes: bytes) -> bytes:
    pdf = documents._libreoffice_convert(docx_bytes, ".docx", "pdf", ".pdf")
    assert pdf, "LibreOffice не смогла собрать тестовый PDF"
    return pdf


def _pdf_info(pdf: bytes) -> tuple[int, int, str]:
    """``(страниц, картинок, текст)`` PDF."""
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(pdf))
    images = sum(len(page.images) for page in reader.pages)
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    return len(reader.pages), images, text


# --- .odg: до абзацев в draw:frame дотягивается _rewrite_odt_paragraphs ------

@_NEEDS_SOFFICE
def test_odt_rewriter_reaches_text_inside_draw_frames():
    odg = documents.pdf_to_odg_bytes(_to_pdf(_docx_with_image()))
    assert odg
    assert _INN in documents._read_odt_bytes(odg)
    out = documents._rewrite_odt_paragraphs(odg, lambda t: t.replace(_INN, "[ИНН_1]"))
    text = documents._read_odt_bytes(out)
    assert _INN not in text and "[ИНН_1]" in text
    with zipfile.ZipFile(io.BytesIO(out)) as z:
        assert any(n.startswith("Pictures/") for n in z.namelist())


# --- полный круг PDF --------------------------------------------------------

@_NEEDS_SOFFICE
def test_pdf_round_trip_keeps_image_and_replaces_text():
    src = _to_pdf(_docx_with_image())
    pages0, images0, text0 = _pdf_info(src)
    assert images0 == 1 and _INN in text0 and "Иванов" in text0

    res = _anonymize("договор.pdf", src)
    assert res["document_name"] == "договор.anon.pdf"
    assert res["document_mime"] == "application/pdf"
    assert res["document_source"] == "converted"
    assert not any("оформление" in w for w in res["warnings"])

    pages1, images1, text1 = _pdf_info(base64.b64decode(res["document_base64"]))
    assert pages1 == pages0
    assert images1 == images0
    assert _INN not in text1 and _EMAIL not in text1
    assert "[" in text1  # плейсхолдеры на месте значений


# --- скан -------------------------------------------------------------------

@_NEEDS_SOFFICE
def test_scanned_pdf_is_refused_and_nothing_is_returned():
    scan = _to_pdf(_docx_only_image())
    assert _pdf_info(scan)[2].strip() == ""
    with pytest.raises(server._ScanRefused) as exc:
        _anonymize("скан.pdf", scan)
    assert str(exc.value) == documents.SCANNED_PDF_MESSAGE
    assert "скан" in str(exc.value) and "OCR" in str(exc.value)


def test_scan_heuristic_thresholds():
    assert documents._looks_scanned([])
    assert documents._looks_scanned([0, 0, 0])
    assert documents._looks_scanned([5] * 10)  # мусорный слой на каждой странице
    assert not documents._looks_scanned([800])
    assert not documents._looks_scanned([300, 0, 400])  # одна пустая страница — не скан
    assert documents._looks_scanned([3000, 0, 0, 0])  # основная часть — картинки


def test_scanned_pdf_is_refused_before_libreoffice(monkeypatch):
    monkeypatch.setattr(documents, "_pdf_page_text_sizes", lambda data: [0])

    def boom(*a, **k):
        raise AssertionError("LibreOffice не должна запускаться для скана")

    monkeypatch.setattr(documents, "_libreoffice_convert", boom)
    with pytest.raises(server._ScanRefused):
        _anonymize("скан.pdf", b"%PDF-1.4 fake")


def test_refusals_share_the_422_base_class():
    assert issubclass(server._ScanRefused, server._DocumentRefused)
    assert issubclass(server._SpecialCategoryRefused, server._DocumentRefused)


# --- нет LibreOffice: документ отдаётся, но с предупреждением ----------------

@pytest.mark.parametrize("filename", ["a.doc", "a.xls", "a.rtf", "a.pdf"])
def test_missing_libreoffice_is_reported_not_silent(monkeypatch, filename):
    monkeypatch.setattr(documents, "_find_soffice", lambda: None)
    monkeypatch.setattr(documents, "_pdf_page_text_sizes", lambda data: [500])
    for reader in ("_read_doc_bytes", "_read_xls_bytes", "_read_rtf_bytes", "_read_pdf_bytes"):
        monkeypatch.setattr(documents, reader, lambda data: f"Почта {_EMAIL}")
    res = _anonymize(filename, b"fake")
    assert base64.b64decode(res["document_base64"])  # документ на руках
    assert res["document_source"] == "text"
    warning = next(w for w in res["warnings"] if "оформление" in w)
    assert "LibreOffice" in warning and "не установлен" in warning
    assert filename[-4:] in warning


@_NEEDS_SOFFICE
def test_pdf_back_conversion_failure_falls_back_to_text_with_warning(monkeypatch):
    src = _to_pdf(_docx_with_image())
    monkeypatch.setattr(documents, "odg_to_pdf_bytes", lambda data: None)
    res = _anonymize("договор.pdf", src)
    assert res["document_name"] == "договор.anon.txt"
    assert res["document_source"] == "text"
    assert any("оформление" in w for w in res["warnings"])
    assert _INN not in base64.b64decode(res["document_base64"]).decode("utf-8")
