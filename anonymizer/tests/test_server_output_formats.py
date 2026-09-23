"""Формат на выходе = формат на входе.

Пользователь загружает .xlsx — получает .xlsx, загружает .odt — получает .odt.
Раньше всё, кроме .docx, вырождалось в .txt, и документ приходилось собирать
заново вручную. Политика живёт в ``documents.prepare_document`` /
``rebuild_document``; здесь она проверяется целиком, через ``_run_anonymize_file``.

Детекторы включены только регулярочные (``DEFAULT_DETECTORS``) — нужен
детерминированный результат без сети и моделей, а адрес почты они ловят
всегда. LibreOffice не запускается: её наличие подменяется на
``documents._libreoffice_convert`` и друзьях.
"""

from __future__ import annotations

import base64
import io
import sys
import tempfile
import zipfile
from contextlib import contextmanager
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from anonymizer import documents, server, usage_log  # noqa: E402
from anonymizer.detectors import DEFAULT_DETECTORS  # noqa: E402

_EMAIL = "ivan.petrov@example.ru"
# Что «LibreOffice» отдаёт на обратном шаге .docx → .rtf в тесте ниже.
_RTF_BACK = b"{\\rtf1 anonymized}"


@contextmanager
def _regex_pipeline():
    """Реальные regex-детекторы, всё остальное выключено."""
    orig = (
        server._DETECTORS,
        server._DEFAULTS,
        server._REVIEW_CFG,
        server._NER_BACKEND,
        server._NEEDS_MODEL_LOCK,
    )
    server._DETECTORS = {"regex": list(DEFAULT_DETECTORS)}
    server._DEFAULTS = {name: False for name in server._STAGE_NAMES}
    server._DEFAULTS["regex"] = True
    server._REVIEW_CFG = None
    server._NER_BACKEND = "none"
    server._NEEDS_MODEL_LOCK = False
    try:
        yield
    finally:
        (
            server._DETECTORS,
            server._DEFAULTS,
            server._REVIEW_CFG,
            server._NER_BACKEND,
            server._NEEDS_MODEL_LOCK,
        ) = orig


@contextmanager
def _temp_usage_log():
    orig = usage_log.LOG_PATH
    with tempfile.TemporaryDirectory() as tmp:
        usage_log.LOG_PATH = Path(tmp) / "usage.jsonl"
        try:
            yield
        finally:
            usage_log.LOG_PATH = orig


def _anonymize(filename: str, raw: bytes) -> dict:
    with _regex_pipeline(), _temp_usage_log():
        return server._run_anonymize_file(
            {"filename": filename, "file_base64": base64.b64encode(raw).decode("ascii")}
        )


def _doc(res: dict) -> bytes:
    return base64.b64decode(res["document_base64"])


@pytest.fixture(autouse=True)
def no_libreoffice(monkeypatch):
    """По умолчанию хоста без LibreOffice: тесты, которым она нужна, включают
    её сами."""
    monkeypatch.setattr(
        documents, "_libreoffice_convert", lambda data, src, to, out: None
    )


# --- форматы, которые правятся напрямую ------------------------------------

def _xlsx_bytes(rows) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_xlsx_comes_back_as_xlsx():
    res = _anonymize("таблица.xlsx", _xlsx_bytes([["Почта", _EMAIL]]))
    assert res["document_name"] == "таблица.anon.xlsx"
    assert res["document_source"] == "original"
    assert res["is_docx"] is True
    # Книга открывается, и адрес в ячейке действительно заменён.
    ws = openpyxl.load_workbook(io.BytesIO(_doc(res))).active
    assert ws.cell(row=1, column=2).value == "[EMAIL_1]"


def test_number_in_a_cell_is_masked_too():
    """Телефон, записанный в ячейку числом, не должен остаться открытым в
    файле только потому, что это не строка: в предпросмотре он замаскирован,
    и расхождения между предпросмотром и документом быть не может."""
    res = _anonymize("таблица.xlsx", _xlsx_bytes([["Телефон", 79161234567]]))
    ws = openpyxl.load_workbook(io.BytesIO(_doc(res))).active
    assert "79161234567" not in str(ws.cell(row=1, column=2).value)


def test_formula_is_left_alone():
    res = _anonymize("таблица.xlsx", _xlsx_bytes([["Итого", "=1+2"], ["Почта", _EMAIL]]))
    ws = openpyxl.load_workbook(io.BytesIO(_doc(res))).active
    assert ws.cell(row=1, column=2).value == "=1+2"


def _odt_bytes(paragraphs_xml: str) -> bytes:
    content = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<office:document-content'
        ' xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"'
        ' xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0">'
        "<office:body><office:text>" + paragraphs_xml + "</office:text></office:body>"
        "</office:document-content>"
    ).encode("utf-8")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr(
            zipfile.ZipInfo("mimetype"),
            b"application/vnd.oasis.opendocument.text",
            zipfile.ZIP_STORED,
        )
        z.writestr("content.xml", content)
    return buf.getvalue()


def test_odt_comes_back_as_odt():
    res = _anonymize("договор.odt", _odt_bytes(f"<text:p>Почта {_EMAIL}</text:p>"))
    assert res["document_name"] == "договор.anon.odt"
    assert res["document_source"] == "original"
    restored = documents.read_text_from_bytes("x.odt", _doc(res))
    assert _EMAIL not in restored and "[EMAIL_1]" in restored


def test_odt_value_split_across_spans_is_still_masked():
    """В ODF значение часто разорвано на несколько text:span — поузловая
    замена его бы не нашла, и в выданном документе адрес остался бы открытым."""
    split = "<text:p>Почта <text:span>ivan.petrov@</text:span>example.ru</text:p>"
    res = _anonymize("договор.odt", _odt_bytes(split))
    restored = documents.read_text_from_bytes("x.odt", _doc(res))
    assert _EMAIL not in restored and "[EMAIL_1]" in restored


def test_odt_stays_a_valid_archive():
    res = _anonymize("договор.odt", _odt_bytes(f"<text:p>Почта {_EMAIL}</text:p>"))
    with zipfile.ZipFile(io.BytesIO(_doc(res))) as z:
        # Требование ODF: mimetype первым и без сжатия.
        assert z.namelist()[0] == "mimetype"
        assert z.getinfo("mimetype").compress_type == zipfile.ZIP_STORED


def test_xml_comes_back_as_xml():
    src = f'<?xml version="1.0" encoding="utf-8"?><doc><mail>{_EMAIL}</mail></doc>'
    res = _anonymize("данные.xml", src.encode("utf-8"))
    assert res["document_name"] == "данные.anon.xml"
    assert res["document_mime"] == "application/xml"
    body = _doc(res).decode("utf-8")
    assert _EMAIL not in body and "[EMAIL_1]" in body
    assert "<mail>" in body  # разметка на месте


# --- простой текст: документ и есть его текст ------------------------------

def test_json_keeps_its_extension():
    res = _anonymize("данные.json", f'{{"mail": "{_EMAIL}"}}'.encode("utf-8"))
    assert res["document_name"] == "данные.anon.json"
    assert res["document_mime"] == "application/json"
    # Простой текст клиент подставляет сам — сервер документ не собирает.
    assert res["is_docx"] is False
    assert _EMAIL not in _doc(res).decode("utf-8")


def test_csv_keeps_its_extension():
    res = _anonymize("список.csv", f"почта;{_EMAIL}".encode("utf-8"))
    assert res["document_name"] == "список.anon.csv"
    assert res["document_mime"] == "text/csv"


# --- форматы, которым нужна конвертация ------------------------------------

def test_xls_without_libreoffice_becomes_a_workbook_not_text(monkeypatch):
    """Запись .xls в Python'е нет вовсе, поэтому старая книга отдаётся .xlsx —
    но именно книгой, а не текстом."""
    monkeypatch.setattr(documents, "_read_xls_bytes", lambda data: f"Почта\t{_EMAIL}")
    res = _anonymize("таблица.xls", b"fake xls")
    assert res["document_name"] == "таблица.anon.xlsx"
    assert res["document_source"] == "text"
    ws = openpyxl.load_workbook(io.BytesIO(_doc(res))).active
    assert ws.cell(row=1, column=2).value == "[EMAIL_1]"


def test_rtf_round_trips_back_to_rtf(monkeypatch):
    """RTF правится кругом rtf → docx → правка → rtf: напрямую значения в нём
    не найти, они разорваны управляющими словами."""
    import docx as docx_mod

    def fake_convert(data, src_suffix, convert_to, out_ext):
        if out_ext == ".docx":
            document = docx_mod.Document()
            document.add_paragraph(f"Почта {_EMAIL}")
            buf = io.BytesIO()
            document.save(buf)
            return buf.getvalue()
        return _RTF_BACK  # обратный шаг

    monkeypatch.setattr(documents, "_libreoffice_convert", fake_convert)
    res = _anonymize("письмо.rtf", b"{\\rtf1 fake}")
    assert res["document_name"] == "письмо.anon.rtf"
    assert res["document_mime"] == "application/rtf"
    assert res["document_source"] == "converted"
    # Отдан результат ОБРАТНОЙ конвертации, а не промежуточный .docx.
    assert _doc(res) == _RTF_BACK


def test_rtf_without_libreoffice_degrades_to_a_word_document(monkeypatch):
    monkeypatch.setattr(documents, "_read_rtf_bytes", lambda data: f"Почта {_EMAIL}")
    res = _anonymize("письмо.rtf", b"{\\rtf1 fake}")
    assert res["document_name"] == "письмо.anon.docx"
    assert res["document_source"] == "text"


# --- PDF: сознательно остаётся текстом -------------------------------------

def test_pdf_stays_text_on_purpose(monkeypatch):
    """Решение принято отдельно: переписать текст в PDF без потери вёрстки
    умеет только PyMuPDF под AGPL-3. Тест фиксирует решение, чтобы PDF не
    «починили» случайно кругом через LibreOffice."""
    monkeypatch.setattr(documents, "_read_pdf_bytes", lambda data: f"Почта {_EMAIL}")
    res = _anonymize("скан.pdf", b"%PDF-1.4 fake")
    assert res["document_name"] == "скан.anon.txt"
    assert res["document_mime"] == "text/plain"
    assert res["document_source"] == "text"
