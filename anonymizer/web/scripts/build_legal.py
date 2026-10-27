"""Сборка публичных юридических документов для сайта.

Оферта (Word) и Политика (Markdown) живут в комплекте документов вне
репозитория — это источник истины. Скрипт превращает их в HTML-фрагменты,
которые страницы /legal/* вставляют как есть, и кладёт рядом PDF для
скачивания. Запускать после каждой правки оферты или Политики:

    python scripts/build_legal.py \
        --offer  "<путь>/6_Оферта_с_поручением.docx" \
        --offer-pdf "<путь>/6_Оферта_с_поручением.pdf" \
        --policy "<путь>/5_Политика_обработки_ПДн_по_факту.md" \
        --policy-pdf "<путь>/5_Политика_обработки_ПДн_по_факту.pdf"

Нумерация пунктов в оферте — автоматическая нумерация Word; python-docx её не
отдаёт, поэтому счётчики восстанавливаются здесь по numbering.xml так же, как
их считает Word: свой счётчик у каждого списка (numId) и уровня, сброс
вложенных уровней при переходе на верхний.

Не забыть: при изменении редакции поднять LEGAL_DOCS_VERSION в lib/legal.ts.
"""
from __future__ import annotations

import argparse
import html
import re
import shutil
import sys
from pathlib import Path

import docx
import markdown
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

WEB = Path(__file__).resolve().parents[1]
OUT_HTML = WEB / "content" / "legal"
OUT_PDF = WEB / "public" / "legal"

# Номер списка, которым в оферте пронумерованы заголовки разделов.
HEADING_NUM_ID = "1"
# Заголовки без автонумерации (реквизиты, приложение с тарифами).
UNNUMBERED_HEADINGS = ("Адреса и реквизиты", "Приложение 1", "ТАРИФЫ")
# Строки, которые имеют смысл только в печатном экземпляре: QR-код в HTML
# не переносится.
SKIP_PREFIXES = ("Воспользуйтесь QR-кодом",)

URL_RE = re.compile(r"(https?://[^\s<>()«»]+[^\s<>()«».,;:])|([\w.+-]+@[\w-]+(?:\.[\w-]+)+)")


def linkify(text: str) -> str:
    out, pos = [], 0
    for m in URL_RE.finditer(text):
        out.append(html.escape(text[pos:m.start()]))
        if m.group(1):
            url = m.group(1)
            out.append(f'<a href="{html.escape(url)}">{html.escape(url)}</a>')
        else:
            mail = m.group(2)
            out.append(f'<a href="mailto:{html.escape(mail)}">{html.escape(mail)}</a>')
        pos = m.end()
    out.append(html.escape(text[pos:]))
    return "".join(out)


class Numbering:
    """Восстановление автоматических номеров Word."""

    def __init__(self, document):
        self.levels: dict[str, dict[int, tuple[str, str, int]]] = {}
        self.counters: dict[str, dict[int, int]] = {}
        part = document.part.numbering_part
        root = part.element
        abstracts = {}
        for a in root.findall(qn("w:abstractNum")):
            lv = {}
            for lvl in a.findall(qn("w:lvl")):
                ilvl = int(lvl.get(qn("w:ilvl")))
                fmt = lvl.find(qn("w:numFmt"))
                txt = lvl.find(qn("w:lvlText"))
                start = lvl.find(qn("w:start"))
                lv[ilvl] = (
                    fmt.get(qn("w:val")) if fmt is not None else "decimal",
                    txt.get(qn("w:val")) if txt is not None else "",
                    int(start.get(qn("w:val"))) if start is not None else 1,
                )
            abstracts[a.get(qn("w:abstractNumId"))] = lv
        for n in root.findall(qn("w:num")):
            num_id = n.get(qn("w:numId"))
            abs_id = n.find(qn("w:abstractNumId")).get(qn("w:val"))
            self.levels[num_id] = dict(abstracts.get(abs_id, {}))

    def label(self, num_id: str, ilvl: int) -> tuple[str, str]:
        """Вернуть (вид, метка): вид 'bullet' или 'number'."""
        lv = self.levels.get(num_id, {})
        fmt, text, start = lv.get(ilvl, ("decimal", "%1.", 1))
        if fmt == "bullet":
            return "bullet", ""
        cnt = self.counters.setdefault(num_id, {})
        cnt[ilvl] = cnt.get(ilvl, start - 1) + 1
        for deeper in [k for k in cnt if k > ilvl]:
            del cnt[deeper]

        def sub(m: re.Match) -> str:
            k = int(m.group(1)) - 1
            return str(cnt.get(k, lv.get(k, ("", "", 1))[2]))

        return "number", re.sub(r"%(\d)", sub, text)


def num_of(p) -> tuple[str, int] | None:
    ppr = p._p.pPr
    if ppr is None or ppr.numPr is None or ppr.numPr.numId is None:
        return None
    num_id = str(ppr.numPr.numId.val)
    if num_id == "0":
        return None
    ilvl = ppr.numPr.ilvl.val if ppr.numPr.ilvl is not None else 0
    return num_id, int(ilvl)


class Writer:
    def __init__(self, numbering: Numbering):
        self.num = numbering
        self.parts: list[str] = []
        self.in_list = False

    def close_list(self):
        if self.in_list:
            self.parts.append("</ul>")
            self.in_list = False

    def paragraph(self, p, *, title_index: int | None = None):
        # Ручные переносы строк внутри абзаца (адрес, заголовок приложения)
        # в вёрстке сайта не нужны — строка переносится сама.
        text = re.sub(r"\s*\n\s*", " ", p.text).strip()
        if not text or any(text.startswith(s) for s in SKIP_PREFIXES):
            return
        n = num_of(p)
        if title_index is not None:
            self.close_list()
            tag = {0: "h1", 1: "p class=\"doc-subtitle\""}.get(title_index, "p class=\"doc-edition\"")
            self.parts.append(f"<{tag}>{linkify(text)}</{tag.split()[0]}>")
            return
        if n is None:
            self.close_list()
            if any(text.startswith(s) for s in UNNUMBERED_HEADINGS):
                self.parts.append(f"<h2>{linkify(text)}</h2>")
            else:
                self.parts.append(f"<p>{linkify(text)}</p>")
            return
        num_id, ilvl = n
        kind, label = self.num.label(num_id, ilvl)
        if kind == "bullet":
            if not self.in_list:
                self.parts.append("<ul>")
                self.in_list = True
            self.parts.append(f"<li>{linkify(text)}</li>")
            return
        self.close_list()
        if num_id == HEADING_NUM_ID and ilvl == 0:
            anchor = label.rstrip(".")
            self.parts.append(f'<h2 id="p{anchor}">{html.escape(label)} {linkify(text)}</h2>')
            return
        self.parts.append(
            f'<p class="clause" id="p{html.escape(label.rstrip("."))}">'
            f'<span class="clause-num">{html.escape(label)}</span> {linkify(text)}</p>'
        )

    def table(self, tbl):
        self.close_list()
        rows = tbl.rows
        ncols = len(tbl.columns)
        seen: list = []

        def unique(row):
            cells = []
            for c in row.cells:
                if any(c._tc is s for s in seen):
                    if cells and cells[-1][0]._tc is c._tc:
                        cells[-1][1] += 1
                    continue
                seen.append(c._tc)
                cells.append([c, 1])
            return cells

        if ncols < 3:
            # Таблица вёрстки (преамбула, реквизиты): содержимое ячеек подряд.
            for row in rows:
                for c, _ in unique(row):
                    for p in c.paragraphs:
                        self.paragraph(p)
                    self.close_list()
            return
        self.parts.append('<div class="table-wrap"><table>')
        for ri, row in enumerate(rows):
            cells = unique(row)
            if not cells:
                continue
            tag = "th" if ri <= 1 else "td"
            self.parts.append("<tr>")
            for c, span in cells:
                inner = "<br>".join(linkify(p.text.strip()) for p in c.paragraphs if p.text.strip())
                colspan = f' colspan="{span}"' if span > 1 else ""
                self.parts.append(f"<{tag}{colspan}>{inner}</{tag}>")
            self.parts.append("</tr>")
        self.parts.append("</table></div>")


def build_offer(path: Path) -> str:
    d = docx.Document(str(path))
    w = Writer(Numbering(d))
    title_seen = 0
    for el in d.element.body:
        tag = el.tag.split("}")[1]
        if tag == "p":
            p = Paragraph(el, d)
            if title_seen < 3 and p.text.strip():
                w.paragraph(p, title_index=title_seen)
                title_seen += 1
            else:
                w.paragraph(p)
        elif tag == "tbl":
            w.table(Table(el, d))
    w.close_list()
    return "\n".join(w.parts)


def build_policy(path: Path) -> str:
    text = path.read_text(encoding="utf-8")
    # Строки «Редакция…» и «Дата вступления в силу…» идут подряд без пустой
    # строки — в Markdown они слились бы в один абзац.
    text = re.sub(r"^(Редакция:[^\n]*)\n(Дата вступления)", r"\1  \n\2", text, flags=re.M)
    body = markdown.markdown(text, extensions=["tables", "sane_lists"], output_format="html")
    body = re.sub(r"<h2>", '<p class="doc-subtitle">', body, count=1).replace("</h2>", "</p>", 1)
    body = body.replace("<table>", '<div class="table-wrap"><table>').replace("</table>", "</table></div>")
    body = re.sub(
        r"(?<![\"/>])(https?://[^\s<>()«»]+[^\s<>()«».,;:])",
        lambda m: f'<a href="{m.group(1)}">{m.group(1)}</a>',
        body,
    )
    return body


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--offer", required=True, type=Path)
    ap.add_argument("--offer-pdf", required=True, type=Path)
    ap.add_argument("--policy", required=True, type=Path)
    ap.add_argument("--policy-pdf", required=True, type=Path)
    a = ap.parse_args()

    OUT_HTML.mkdir(parents=True, exist_ok=True)
    OUT_PDF.mkdir(parents=True, exist_ok=True)
    (OUT_HTML / "offer.html").write_text(build_offer(a.offer), encoding="utf-8")
    (OUT_HTML / "privacy.html").write_text(build_policy(a.policy), encoding="utf-8")
    shutil.copyfile(a.offer_pdf, OUT_PDF / "oferta.pdf")
    shutil.copyfile(a.policy_pdf, OUT_PDF / "politika.pdf")
    print("готово:", OUT_HTML, OUT_PDF)


if __name__ == "__main__":
    main()
